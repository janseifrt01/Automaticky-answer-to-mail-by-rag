"""Tests for file parsers (Task 3.1)."""

from __future__ import annotations

import io

import pytest
from app.rag import parsers
from docx import Document


def test_extract_txt():
    assert parsers.extract_text("notes.txt", b"hello world") == "hello world"


def test_extract_md_uses_plain_reader():
    assert parsers.extract_text("readme.md", b"# Title\nbody") == "# Title\nbody"


def test_extract_docx():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("First line")
    doc.add_paragraph("Second line")
    doc.save(buf)
    text = parsers.extract_text("doc.docx", buf.getvalue())
    assert "First line" in text and "Second line" in text


def test_unsupported_type_raises():
    with pytest.raises(parsers.UnsupportedFileType):
        parsers.extract_text("image.png", b"\x89PNG")
